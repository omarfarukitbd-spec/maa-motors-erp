import docx
from docx.shared import Pt, RGBColor, Mm
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from bijoy2unicode import converter
import subprocess
import pymupdf
import os

conv = converter.Unicode()

doc = docx.Document()
p = doc.add_paragraph()
r = p.add_run(conv.convertUnicodeToBijoy('বিসমিল্লাহির রাহমানির রাহিম'))
r.font.name = 'SutonnyMJ'
r.font.size = Pt(12)
r.bold = True
rPr = r._r.get_or_add_rPr()
rPr.append(parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="SutonnyMJ" w:hAnsi="SutonnyMJ"/>'))

p2 = doc.add_paragraph()
r2 = p2.add_run(conv.convertUnicodeToBijoy('ইশরাক টাওয়ার'))
r2.font.name = 'SutonnyMJ'
r2.font.size = Pt(24)
r2.bold = True
r2.font.color.rgb = RGBColor(123, 17, 19)
rPr2 = r2._r.get_or_add_rPr()
rPr2.append(parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="SutonnyMJ" w:hAnsi="SutonnyMJ"/>'))

p3 = doc.add_paragraph()
r3 = p3.add_run(conv.convertUnicodeToBijoy('মালিক: মোহাম্মদ এমরান'))
r3.font.name = 'SutonnyMJ'
r3.font.size = Pt(12)
r3.bold = True
rPr3 = r3._r.get_or_add_rPr()
rPr3.append(parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="SutonnyMJ" w:hAnsi="SutonnyMJ"/>'))

docx_p = os.path.abspath('test_sutonny.docx')
pdf_p = os.path.abspath('test_sutonny.pdf')
doc.save(docx_p)

cmd = ['powershell', '-ExecutionPolicy', 'Bypass', '-File', 'E:\\maa-motors-erp\\convert_to_pdf.ps1', '-docxPath', docx_p, '-pdfPath', pdf_p]
subprocess.run(cmd, capture_output=True, text=True)

doc_pdf = pymupdf.open(pdf_p)
pix = doc_pdf[0].get_pixmap(dpi=150)
pix.save('test_sutonny_preview.png')
print('Rendered test_sutonny_preview.png')
