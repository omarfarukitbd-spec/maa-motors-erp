import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    """Set inner padding for table cells in dxa (1 pt = 20 dxa)"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_background(cell, fill_hex):
    """Set background color of a cell (e.g. '0284C7')"""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_borders(cell, top="CBD5E1", bottom="CBD5E1", left="CBD5E1", right="CBD5E1", sz="4", val="single"):
    """Set individual borders for a table cell"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="{val if top else "none"}" w:sz="{sz}" w:space="0" w:color="{top if top else "auto"}"/>'
        f'<w:bottom w:val="{val if bottom else "none"}" w:sz="{sz}" w:space="0" w:color="{bottom if bottom else "auto"}"/>'
        f'<w:left w:val="{val if left else "none"}" w:sz="{sz}" w:space="0" w:color="{left if left else "auto"}"/>'
        f'<w:right w:val="{val if right else "none"}" w:sz="{sz}" w:space="0" w:color="{right if right else "auto"}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

def add_run(paragraph, text, font_name="Segoe UI", cs_font="Kalpurush", size_pt=9.5, bold=False, italic=False, color=None):
    """Add a run with both Latin and Complex Script (Bengali) font support"""
    run = paragraph.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:cs'), cs_font)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    return run

def create_container_docx(output_path):
    doc = docx.Document()
    logo_path = r"E:\maa-motors-erp\Web_ERP\public\shop-official-logo.jpg"
    
    # 1. Page Setup (Exact A4, Margins: Top/Bottom 0.35 in, Left/Right 0.45 in)
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.35)
    section.bottom_margin = Inches(0.35)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)
    
    # Theme Colors
    OCEAN_BLUE = RGBColor(2, 132, 199)   # #0284C7
    DEEP_BLUE = RGBColor(3, 105, 161)    # #0369A1
    WHITE = RGBColor(255, 255, 255)
    DARK_SLATE = RGBColor(30, 41, 59)    # #1E293B
    MUTED_GRAY = RGBColor(71, 85, 105)   # #475569
    RED = RGBColor(220, 38, 38)          # #DC2626
    EMERALD = RGBColor(22, 101, 52)      # #166534
    
    # ----------------------------------------------------
    # 1. Official App Header Table (Blue Gradient Style)
    # ----------------------------------------------------
    hdr_table = doc.add_table(rows=1, cols=3)
    hdr_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_table.autofit = False
    
    # Cell 0: Official Logo Box
    c_logo = hdr_table.cell(0, 0)
    c_logo.width = Inches(1.1)
    set_cell_margins(c_logo, top=100, bottom=100, left=100, right=60)
    set_cell_background(c_logo, "0284C7")
    set_cell_borders(c_logo, top=None, bottom=None, left=None, right=None)
    p_logo = c_logo.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists(logo_path):
        p_logo.add_run().add_picture(logo_path, width=Inches(0.85))
        
    # Cell 1: Shop Information
    c_info = hdr_table.cell(0, 1)
    c_info.width = Inches(4.35)
    set_cell_margins(c_info, top=100, bottom=100, left=80, right=80)
    set_cell_background(c_info, "0284C7")
    set_cell_borders(c_info, top=None, bottom=None, left=None, right=None)
    p_info = c_info.paragraphs[0]
    p_info.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_info.paragraph_format.space_before = Pt(0)
    p_info.paragraph_format.space_after = Pt(0)
    
    add_run(p_info, "M/S. MAA-MOTOR'S\n", size_pt=16, bold=True, color=WHITE)
    add_run(p_info, "Proprietor: Mohammed Amran | প্রোঃ মোঃ মোহাম্মদ এমরান\n", size_pt=8.5, bold=True, color=WHITE)
    add_run(p_info, "আমদানিকারক ও পাইকারী বিক্রেতা — দোকান নং-২২, রহমান টাওয়ার, ১ম রেল গেইট, মুরাদপুর, হাটহাজারী রোড, চট্টগ্রাম\n", size_pt=7.5, color=WHITE)
    add_run(p_info, "Mobile: 01819-397669, 01815-707934, 01818-195690", size_pt=8, bold=True, color=WHITE)
    
    # Cell 2: Title Badge
    c_badge = hdr_table.cell(0, 2)
    c_badge.width = Inches(1.92)
    set_cell_margins(c_badge, top=100, bottom=100, left=60, right=100)
    set_cell_background(c_badge, "0369A1")
    set_cell_borders(c_badge, top=None, bottom=None, left=None, right=None)
    p_badge = c_badge.paragraphs[0]
    p_badge.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_badge.paragraph_format.space_before = Pt(0)
    p_badge.paragraph_format.space_after = Pt(0)
    
    add_run(p_badge, "CONTAINER REPORT\n", size_pt=11, bold=True, color=WHITE)
    add_run(p_badge, "কন্টেইনার চালান ও খালাস খরচ বিবরণী", size_pt=8, bold=True, color=WHITE)
    
    p_sp1 = doc.add_paragraph()
    p_sp1.paragraph_format.space_before = Pt(0)
    p_sp1.paragraph_format.space_after = Pt(3)

    # ----------------------------------------------------
    # 2. Sub-bar: Bismillah, Chalan Check & Delivery Party
    # ----------------------------------------------------
    sub_table = doc.add_table(rows=1, cols=3)
    sub_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sub_table.autofit = False
    
    col_w_sub = [Inches(3.2), Inches(1.6), Inches(2.57)]
    for c_idx, cell in enumerate(sub_table.rows[0].cells):
        cell.width = col_w_sub[c_idx]
        set_cell_margins(cell, top=40, bottom=40, left=80, right=80)
        set_cell_background(cell, "F8FAFC")
        set_cell_borders(cell, top="CBD5E1", bottom="CBD5E1", left="CBD5E1", right="CBD5E1", sz="4")
        
    p_s0 = sub_table.rows[0].cells[0].paragraphs[0]
    p_s0.paragraph_format.space_before = Pt(0)
    p_s0.paragraph_format.space_after = Pt(0)
    add_run(p_s0, "بِسْمِ اللَّهِ الرَّحْمَٰنِ الرَّحِيمِ (বিসমিল্লাহির রাহমানির রাহিম)", size_pt=8, italic=True, color=MUTED_GRAY)
    
    p_s1 = sub_table.rows[0].cells[1].paragraphs[0]
    p_s1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_s1.paragraph_format.space_before = Pt(0)
    p_s1.paragraph_format.space_after = Pt(0)
    add_run(p_s1, "কোটেশন [  ]  চালান [ ✓ ]", size_pt=8.5, bold=True, color=DEEP_BLUE)
    
    p_s2 = sub_table.rows[0].cells[2].paragraphs[0]
    p_s2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_s2.paragraph_format.space_before = Pt(0)
    p_s2.paragraph_format.space_after = Pt(0)
    add_run(p_s2, "ডেলিভারি পার্টি: .......................................", size_pt=8.5, bold=True, color=DARK_SLATE)

    p_sp2 = doc.add_paragraph()
    p_sp2.paragraph_format.space_before = Pt(0)
    p_sp2.paragraph_format.space_after = Pt(3)

    # ----------------------------------------------------
    # 3. Meta Table (4 Columns)
    # ----------------------------------------------------
    meta_table = doc.add_table(rows=4, cols=4)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    
    meta_data = [
        [("কন্টেইনার নং:", True), ("", False), ("কন্টেইনার সাইজ:", True), ("২০ ফুট [  ]   ৪০ ফুট [ ✓ ]", False)],
        [("এল/সি নং (L/C No):", True), ("", False), ("ইনভয়েস মূল্য (U$D):", True), ("$ ", False)],
        [("লোডিং তারিখ:", True), ("...... / ...... / ২০......", False), ("ডেলিভারি তারিখ:", True), ("...... / ...... / ২০......", False)],
        [("কোম্পানী / শিপার:", True), ("", False), ("আগমন বন্দর / ডিপো:", True), ("চট্টগ্রাম বন্দর [  ]   কমলাপুর আইসিডি [ ✓ ]", False)]
    ]
    
    col_widths_meta = [Inches(1.8), Inches(1.88), Inches(1.8), Inches(1.89)]
    for r_idx, row in enumerate(meta_table.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.width = col_widths_meta[c_idx]
            text, is_label = meta_data[r_idx][c_idx]
            set_cell_margins(cell, top=50, bottom=50, left=80, right=80)
            set_cell_borders(cell, top="CBD5E1", bottom="CBD5E1", left="CBD5E1", right="CBD5E1", sz="4")
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            
            if is_label:
                set_cell_background(cell, "F0F9FF")
                add_run(p, text, size_pt=8, bold=True, color=DEEP_BLUE)
            else:
                add_run(p, text, size_pt=8, color=DARK_SLATE)

    p_sp3 = doc.add_paragraph()
    p_sp3.paragraph_format.space_before = Pt(0)
    p_sp3.paragraph_format.space_after = Pt(3)

    # ----------------------------------------------------
    # 4. Expense Items Table
    # ----------------------------------------------------
    expense_items = [
        ("১", "UAE প্রাথমিক খরচ (বকশিশ, চা-নাস্তা, স্কেল / ওয়েব্রিজ ফি ও এডভান্স/AD সমন্বয়)"),
        ("২", "UAE লোডিং, শিপিং লাইন চার্জ, L/C কমিশন ও বৈদেশিক ব্যাংক পেমেন্ট"),
        ("৩", "ব্যাংক চার্জ ও মার্জিন (ইসলামী ব্যাংক L/C মার্জিন / আরিয়ান বা অন্যান্য ফি)"),
        ("৪", "কাস্টমস ক্লিয়ারেন্স সার্টিফিকেট ফি (খালেক বা নাসির এজেন্ট)"),
        ("৫", "পোর্ট ইয়ার্ড কন্টেইনার মুভমেন্ট / ঠেলা ও গেট পারমিশন চার্জ"),
        ("৬", "কাস্টমস পরীক্ষণ (Customs Examination) ও ঢাকা যাতায়াত খরচ"),
        ("৭", "গোডাউন যাতায়াত ও খালাস তদারকি চা-নাস্তা খরচ"),
        ("৮", "লোকাল পরিবহন কাভার্ড ভ্যান ভাড়া, খালাস লেবার ও ক্রেন চার্জ"),
        ("৯", "অস্থায়ী গোডাউন ভাড়া (আম্বিয়া আনিস + জাহেদ আনিস স্পেস বাবদ)"),
        ("১০", "C&F কমিশন, বন্দর মাশুল (Port Bill), শিপিং বিল, ডক লেবার বিল ও হ্যান্ডলিং চার্জ"),
        ("১১", "ঘাট খেয়া পারাপার ও চ্যানেল পরিবহন খরচ"),
        ("১২", "বিবিধ আনুষঙ্গিক খরচ (অন্যান্য)"),
        ("১৩", "অতিরিক্ত খাত (প্রয়োজন অনুযায়ী কাস্টম এন্ট্রি)")
    ]
    
    exp_table = doc.add_table(rows=1 + len(expense_items) + 3, cols=4)
    exp_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    exp_table.autofit = False
    
    tbl_widths = [Inches(0.42), Inches(4.55), Inches(1.25), Inches(1.15)]
    
    # 4a. Header Row
    headers = ["ক্র.", "খরচের খাত ও কাজের বিবরণ (Expense Particulars)", "পরিমাণ (টাকা)", "মন্তব্য / ভাউচার"]
    for c_idx, h_text in enumerate(headers):
        cell = exp_table.cell(0, c_idx)
        cell.width = tbl_widths[c_idx]
        set_cell_margins(cell, top=70, bottom=70, left=80, right=80)
        set_cell_background(cell, "0369A1")
        set_cell_borders(cell, top="0369A1", bottom="0369A1", left="0369A1", right="0369A1", sz="6")
        
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 2, 3] else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        add_run(p, h_text, size_pt=8, bold=True, color=WHITE)
        
    # 4b. Items
    for idx, (sl, desc) in enumerate(expense_items):
        row_idx = 1 + idx
        row_cells = exp_table.rows[row_idx].cells
        bg_hex = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
        
        # Col 0: SL
        cell_sl = row_cells[0]
        cell_sl.width = tbl_widths[0]
        set_cell_margins(cell_sl, top=35, bottom=35, left=50, right=50)
        set_cell_background(cell_sl, bg_hex)
        set_cell_borders(cell_sl, top="CBD5E1", bottom="CBD5E1", left="CBD5E1", right="CBD5E1", sz="4")
        p_sl = cell_sl.paragraphs[0]
        p_sl.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sl.paragraph_format.space_before = Pt(0)
        p_sl.paragraph_format.space_after = Pt(0)
        add_run(p_sl, sl, size_pt=7.5, bold=True, color=DEEP_BLUE)
        
        # Col 1: Description
        cell_desc = row_cells[1]
        cell_desc.width = tbl_widths[1]
        set_cell_margins(cell_desc, top=35, bottom=35, left=80, right=80)
        set_cell_background(cell_desc, bg_hex)
        set_cell_borders(cell_desc, top="CBD5E1", bottom="CBD5E1", left="CBD5E1", right="CBD5E1", sz="4")
        p_desc = cell_desc.paragraphs[0]
        p_desc.paragraph_format.space_before = Pt(0)
        p_desc.paragraph_format.space_after = Pt(0)
        add_run(p_desc, desc, size_pt=7.5, color=DARK_SLATE)
        
        # Col 2: Amount
        cell_amt = row_cells[2]
        cell_amt.width = tbl_widths[2]
        set_cell_margins(cell_amt, top=35, bottom=35, left=80, right=80)
        set_cell_background(cell_amt, bg_hex)
        set_cell_borders(cell_amt, top="CBD5E1", bottom="CBD5E1", left="CBD5E1", right="CBD5E1", sz="4")
        p_amt = cell_amt.paragraphs[0]
        p_amt.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_amt.paragraph_format.space_before = Pt(0)
        p_amt.paragraph_format.space_after = Pt(0)
        add_run(p_amt, "", size_pt=7.5, color=DARK_SLATE)
        
        # Col 3: Remarks
        cell_rem = row_cells[3]
        cell_rem.width = tbl_widths[3]
        set_cell_margins(cell_rem, top=35, bottom=35, left=50, right=50)
        set_cell_background(cell_rem, bg_hex)
        set_cell_borders(cell_rem, top="CBD5E1", bottom="CBD5E1", left="CBD5E1", right="CBD5E1", sz="4")
        p_rem = cell_rem.paragraphs[0]
        p_rem.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_rem.paragraph_format.space_before = Pt(0)
        p_rem.paragraph_format.space_after = Pt(0)
        add_run(p_rem, "", size_pt=7.5, color=DARK_SLATE)

    # 4c. Calculation Rows
    curr_row = 1 + len(expense_items)
    
    # Total
    r_total = exp_table.rows[curr_row].cells
    r_total[0].merge(r_total[1])
    set_cell_margins(r_total[0], top=50, bottom=50, left=80, right=80)
    set_cell_background(r_total[0], "E0F2FE")
    set_cell_borders(r_total[0], top="0369A1", bottom="94A3B8", left="CBD5E1", right="CBD5E1", sz="6")
    p_tot_lbl = r_total[0].paragraphs[0]
    p_tot_lbl.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_tot_lbl.paragraph_format.space_before = Pt(0)
    p_tot_lbl.paragraph_format.space_after = Pt(0)
    add_run(p_tot_lbl, "প্রতি কন্টেইনার সর্বমোট খালাস খরচ (TOTAL EXPENSES):", size_pt=8, bold=True, color=DEEP_BLUE)
    
    set_cell_margins(r_total[2], top=50, bottom=50, left=80, right=80)
    set_cell_background(r_total[2], "E0F2FE")
    set_cell_borders(r_total[2], top="0369A1", bottom="94A3B8", left="CBD5E1", right="CBD5E1", sz="6")
    p_tot_val = r_total[2].paragraphs[0]
    p_tot_val.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_tot_val.paragraph_format.space_before = Pt(0)
    p_tot_val.paragraph_format.space_after = Pt(0)
    add_run(p_tot_val, "= TK", size_pt=8, bold=True, color=DEEP_BLUE)
    
    set_cell_margins(r_total[3], top=50, bottom=50, left=50, right=50)
    set_cell_background(r_total[3], "E0F2FE")
    set_cell_borders(r_total[3], top="0369A1", bottom="94A3B8", left="CBD5E1", right="CBD5E1", sz="6")
    p_tot_rem = r_total[3].paragraphs[0]
    p_tot_rem.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p_tot_rem, "সর্বমোট ব্যয়", size_pt=7, bold=True, color=DEEP_BLUE)
    
    # Advance
    curr_row += 1
    r_adv = exp_table.rows[curr_row].cells
    r_adv[0].merge(r_adv[1])
    set_cell_margins(r_adv[0], top=50, bottom=50, left=80, right=80)
    set_cell_background(r_adv[0], "F0FDF4")
    set_cell_borders(r_adv[0], top="94A3B8", bottom="94A3B8", left="CBD5E1", right="CBD5E1", sz="4")
    p_adv_lbl = r_adv[0].paragraphs[0]
    p_adv_lbl.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_adv_lbl.paragraph_format.space_before = Pt(0)
    p_adv_lbl.paragraph_format.space_after = Pt(0)
    add_run(p_adv_lbl, "বাদ: প্রাপ্ত নগদ / ব্যাংক অগ্রিম জমা (Less: Advance Received):", size_pt=8, bold=True, color=EMERALD)
    
    set_cell_margins(r_adv[2], top=50, bottom=50, left=80, right=80)
    set_cell_background(r_adv[2], "F0FDF4")
    set_cell_borders(r_adv[2], top="94A3B8", bottom="94A3B8", left="CBD5E1", right="CBD5E1", sz="4")
    p_adv_val = r_adv[2].paragraphs[0]
    p_adv_val.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_run(p_adv_val, "", size_pt=8, bold=True, color=EMERALD)
    
    set_cell_margins(r_adv[3], top=50, bottom=50, left=50, right=50)
    set_cell_background(r_adv[3], "F0FDF4")
    set_cell_borders(r_adv[3], top="94A3B8", bottom="94A3B8", left="CBD5E1", right="CBD5E1", sz="4")
    p_adv_rem = r_adv[3].paragraphs[0]
    p_adv_rem.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p_adv_rem, "জমা / ক্রেডিট (+)", size_pt=7, bold=True, color=EMERALD)
    
    # Net Balance Due
    curr_row += 1
    r_net = exp_table.rows[curr_row].cells
    r_net[0].merge(r_net[1])
    set_cell_margins(r_net[0], top=55, bottom=55, left=80, right=80)
    set_cell_background(r_net[0], "FEF2F2")
    set_cell_borders(r_net[0], top="94A3B8", bottom="DC2626", left="CBD5E1", right="CBD5E1", sz="6")
    p_net_lbl = r_net[0].paragraphs[0]
    p_net_lbl.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_net_lbl.paragraph_format.space_before = Pt(0)
    p_net_lbl.paragraph_format.space_after = Pt(0)
    add_run(p_net_lbl, "সর্বমোট নিট অবশিষ্ট বকেয়া / ব্যালেন্স (TOTAL NET BALANCE DUE):", size_pt=8.5, bold=True, color=RED)
    
    set_cell_margins(r_net[2], top=55, bottom=55, left=80, right=80)
    set_cell_background(r_net[2], "FEF2F2")
    set_cell_borders(r_net[2], top="94A3B8", bottom="DC2626", left="CBD5E1", right="CBD5E1", sz="6")
    p_net_val = r_net[2].paragraphs[0]
    p_net_val.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_run(p_net_val, "= TK", size_pt=8.5, bold=True, color=RED)
    
    set_cell_margins(r_net[3], top=55, bottom=55, left=50, right=50)
    set_cell_background(r_net[3], "FEF2F2")
    set_cell_borders(r_net[3], top="94A3B8", bottom="DC2626", left="CBD5E1", right="CBD5E1", sz="6")
    p_net_rem = r_net[3].paragraphs[0]
    p_net_rem.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p_net_rem, "নিট বকেয়া (-)", size_pt=7, bold=True, color=RED)

    # ----------------------------------------------------
    # 5. In Words
    # ----------------------------------------------------
    p_words = doc.add_paragraph()
    p_words.paragraph_format.space_before = Pt(4)
    p_words.paragraph_format.space_after = Pt(8)
    add_run(p_words, "কথায় (In Words): ", size_pt=8, bold=True, color=DEEP_BLUE)
    add_run(p_words, "................................................................................................................................................................... টাকা মাত্র।", size_pt=8, color=MUTED_GRAY)

    # ----------------------------------------------------
    # 6. Signatures
    # ----------------------------------------------------
    sig_table = doc.add_table(rows=1, cols=3)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.autofit = False
    
    sig_widths = [Inches(2.45), Inches(2.45), Inches(2.47)]
    for c_idx, cell in enumerate(sig_table.rows[0].cells):
        cell.width = sig_widths[c_idx]
        set_cell_borders(cell, top=None, bottom=None, left=None, right=None)
        
    p_s1 = sig_table.rows[0].cells[0].paragraphs[0]
    p_s1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p_s1, "...................................................\n", size_pt=8, color=MUTED_GRAY)
    add_run(p_s1, "প্রস্তুতকারকের স্বাক্ষর\n", size_pt=8, bold=True, color=DEEP_BLUE)
    add_run(p_s1, "(Prepared By)", size_pt=7, italic=True, color=MUTED_GRAY)
    
    p_s2 = sig_table.rows[0].cells[1].paragraphs[0]
    p_s2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p_s2, "...................................................\n", size_pt=8, color=MUTED_GRAY)
    add_run(p_s2, "হিসাবরক্ষক / যাচাইকারী\n", size_pt=8, bold=True, color=DEEP_BLUE)
    add_run(p_s2, "(Verified & Checked By)", size_pt=7, italic=True, color=MUTED_GRAY)
    
    p_s3 = sig_table.rows[0].cells[2].paragraphs[0]
    p_s3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p_s3, "...................................................\n", size_pt=8, color=MUTED_GRAY)
    add_run(p_s3, "মেসার্স মা মটরস্\n", size_pt=8.5, bold=True, color=DEEP_BLUE)
    add_run(p_s3, "মোঃ এমরান (আইডি-৯৯২৯)\n", size_pt=7.5, bold=True, color=DARK_SLATE)
    add_run(p_s3, "স্বত্বাধিকারী / অনুমোদিত স্বাক্ষর", size_pt=7, italic=True, color=MUTED_GRAY)
    
    doc.save(output_path)
    print(f"Updated Word Doc: {output_path}")

if __name__ == "__main__":
    out_file = r"E:\maa-motors-erp\Container_Expense_Statement_Form.docx"
    create_container_docx(out_file)
